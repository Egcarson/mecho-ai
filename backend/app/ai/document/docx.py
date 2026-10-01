from app.ai.document.base import DocumentExtractor
from docx import Document


class DOCXExtractor(DocumentExtractor):

    async def extract(
        self,
        file_path: str,
    ) -> str:

        document = Document(
            file_path,
        )

        parts: list[str] = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()

            if text:
                parts.append(text)

        for table in document.tables:
            for row in table.rows:
                cells = [
                    cell.text.strip()
                    for cell in row.cells
                    if cell.text.strip()
                ]

                if cells:
                    parts.append(
                        " | ".join(cells)
                    )

        return "\n\n".join(parts)